from pathlib import Path
from datetime import datetime
import re
import json

import fitz
import pdfplumber
from unidecode import unidecode
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

FUENTE_CONTENIDO = "Arial Narrow"
TAMANO_CONTENIDO = 8

ENCABEZADOS = {
    "informacion_general": ["INFORMACION GENERAL"],
    "presentacion": ["PRESENTACION"],
    "continuidad": ["CONTINUIDAD CURRICULAR"],
    "competencias": ["COMPETENCIAS", "COMPETENCIAS Y AFIRMACIONES ASOCIADAS"],
    "resultados": ["RESULTADOS DE APRENDIZAJE"],
    "metodologia": ["METODOLOGIA"],
    "tematicas": ["TEMATICAS O CONTENIDOS", "CONTENIDOS", "CRONOGRAMA"],
    "evaluacion": ["EVALUACION"],
    "bibliografia": ["REFERENCIAS", "BIBLIOGRAFIA", "REFERENCIAS BIBLIOGRAFICAS"],
}
ORDEN = list(ENCABEZADOS.keys())

MESES = r"ene|feb|mar|abr|may|jun|jul|ago|sep|sept|oct|nov|dic"
DATE_PREFIX_RE = re.compile(rf"^\s*(?:\d{{1,2}}\s+(?:{MESES})\.?)\s+", flags=re.I)


def limpiar_basura(texto: str) -> str:
    if not texto:
        return ""
    for b in ["￾", "\ufeff", "\ufffe", "\x00"]:
        texto = texto.replace(b, "")
    return texto


def corregir_acentos(texto: str) -> str:
    if not texto:
        return ""
    reemplazos = {
        "´A": "Á", "´E": "É", "´I": "Í", "´O": "Ó", "´U": "Ú",
        "´a": "á", "´e": "é", "´i": "í", "´o": "ó", "´u": "ú",
        "A´": "Á", "E´": "É", "I´": "Í", "O´": "Ó", "U´": "Ú",
        "a´": "á", "e´": "é", "i´": "í", "o´": "ó", "u´": "ú",
        "˜n": "ñ", "˜N": "Ñ", "Dise˜nar": "Diseñar", "dise˜nar": "diseñar",
        "dise˜no": "diseño", "disen˜o": "diseño", "desempe˜no": "desempeño",
        "acompa˜na": "acompaña", "tama˜no": "tamaño", "tama˜nos": "tamaños",
        "peque˜nos": "pequeños", "ense˜nanza": "enseñanza",
        "Matem´aticas": "Matemáticas", "Programaci´on": "Programación",
        "Informaci´on": "Información", "Presentaci´on": "Presentación",
        "Evaluaci´on": "Evaluación", "Metodolog´ia": "Metodología", "Metodolog´ıa": "Metodología",
        "Bibliograf´ia": "Bibliografía", "Bibliograf´ıa": "Bibliografía", "Bibliografa": "Bibliografía",
        "Aut´onomo": "Autónomo", "Cr´editos": "Créditos", "Sesi´on": "Sesión",
        "Tem´aticas": "Temáticas", "´Algebra": "Álgebra", "Alge- bra": "Álgebra",
        "anal´ıtica": "analítica", "l´ınea": "línea", "l´ıneas": "líneas",
        "estad´ıstica": "estadística", "estad´ısticos": "estadísticos", "estad´ısticas": "estadísticas",
        "emp´ırica": "empírica", "cr´ıtico": "crítico", "cr´ıtica": "crítica",
        "t´ecnico": "técnico", "t´ecnica": "técnica", "t´ecnicas": "técnicas", "t´ecnicos": "técnicos",
        "diagn´osti": "diagnósti", "din´amica": "dinámica", "din´amicas": "dinámicas",
        "p´erdida": "pérdida", "p´erdidas": "pérdidas", "b´asica": "básica", "b´asico": "básico",
        "b´asicas": "básicas", "b´asicos": "básicos", "S´ıntesis": "Síntesis", "s´ıntesis": "síntesis",
        "caracter´ısticas": "características", "caracter´ıstica": "característica",
        "anomal´ıas": "anomalías", "intuici´on": "intuición", "visi´on": "visión",
        "pr´actica": "práctica", "pr´acticas": "prácticas", "pr´actico": "práctico", "pr´acticos": "prácticos",
        "arquitect´onicas": "arquitectónicas", "hiperpar´ametros": "hiperparámetros",
        "transformaci´on": "transformación", "regularizaci´on": "regularización", "optimizaci´on": "optimización",
        "activaci´on": "activación", "reconstrucci´on": "reconstrucción", "validaci´on": "validación",
        "atenci´on": "atención", "representaci´on": "representación", "explosi´on": "explosión",
        "resoluci´on": "resolución", "comparaci´on": "comparación", "aplicaci´on": "aplicación",
        "tokenizaci´on": "tokenización", "clasificaci´on": "clasificación", "regresi´on": "regresión",
        "desvanecimiento": "desvanecimiento", "perdida": "pérdida"
    }
    for a, b in reemplazos.items():
        texto = texto.replace(a, b)
    # Une artefactos de flechas extraídas sin espacio.
    texto = texto.replace("→ML", "→ ML").replace("→MLOps", "→ MLOps")
    return texto


def limpiar_texto(texto: str) -> str:
    texto = corregir_acentos(limpiar_basura(texto or ""))
    texto = texto.replace("\u00ad", "")
    texto = re.sub(r"[ \t]+", " ", texto)
    return texto.strip()


def norm(texto: str) -> str:
    texto = corregir_acentos(texto or "")
    texto = unidecode(texto).upper()
    texto = re.sub(r"[^A-Z0-9]+", " ", texto)
    return texto.strip()


def leer_pdf(path_pdf):
    doc = fitz.open(path_pdf)
    paginas = []
    for i, page in enumerate(doc, start=1):
        txt = limpiar_texto(page.get_text("text"))
        paginas.append(f"\n--- PAGINA {i} ---\n{txt}")
    doc.close()
    return "\n".join(paginas)


def lineas_utiles(texto):
    return [limpiar_texto(l) for l in texto.splitlines() if limpiar_texto(l)]


def detectar_encabezados(texto):
    lineas = texto.splitlines()
    encontrados = []
    for i, linea in enumerate(lineas):
        actual = norm(linea)
        combinado = actual
        separado = False
        if re.fullmatch(r"\d+", actual) and i + 1 < len(lineas):
            combinado = actual + " " + norm(lineas[i + 1])
            separado = True
        for clave, opciones in ENCABEZADOS.items():
            for opcion in opciones:
                op = norm(opcion)
                if combinado == op or re.fullmatch(r"\d+\s+" + re.escape(op), combinado):
                    heading_start = i
                    content_start = i + 2 if separado else i + 1
                    encontrados.append((heading_start, content_start, clave))
                    break
    depurados = []
    for item in sorted(encontrados, key=lambda x: x[0]):
        if not depurados or not (depurados[-1][2] == item[2] and abs(depurados[-1][0] - item[0]) <= 2):
            depurados.append(item)
    return depurados, lineas


def extraer_bloques(texto):
    enc, lineas = detectar_encabezados(texto)
    bloques = {k: "" for k in ORDEN}
    for pos, (heading_start, content_start, clave) in enumerate(enc):
        inicio = content_start
        fin = enc[pos + 1][0] if pos + 1 < len(enc) else len(lineas)
        bloques[clave] = limpiar_texto("\n".join(lineas[inicio:fin]))
    return bloques, enc


def extraer_valor_campo(bloque, etiqueta, etiquetas):
    ls = lineas_utiles(bloque)
    et = norm(etiqueta)
    etiquetas_n = [norm(x) for x in etiquetas]
    for i, linea in enumerate(ls):
        nl = norm(linea)
        if nl == et or nl.startswith(et + " "):
            val = re.sub(re.escape(etiqueta), "", linea, flags=re.I).replace(":", "").strip()
            if val and norm(val) != et:
                return limpiar_texto(val)
            acum = []
            for j in range(i + 1, min(i + 6, len(ls))):
                nj = norm(ls[j])
                if any(nj == e or nj.startswith(e + " ") for e in etiquetas_n):
                    break
                acum.append(ls[j])
            return limpiar_texto(" ".join(acum))
    return ""


def extraer_horas(lineas, etiqueta):
    for i, l in enumerate(lineas):
        if etiqueta in norm(l):
            nums = re.findall(r"\d+", l)
            if nums:
                return nums[-1]
            for j in range(i + 1, min(i + 8, len(lineas))):
                if re.search(r"\d+", lineas[j]):
                    return re.findall(r"\d+", lineas[j])[0]
    return ""


def extraer_info_general(bloque_info):
    etiquetas = [
        "PROGRAMA", "ÁREA", "AREA", "ASIGNATURA", "CRÉDITOS", "CREDITOS",
        "SEMESTRE", "HORAS PRESENCIALES", "HORAS PRESENCIA- LES",
        "HORAS DE TRABAJO", "AUTÓNOMO", "AUTONOMO", "PROFESOR", "PROFESORES", "CORREO"
    ]
    ls = lineas_utiles(bloque_info)
    info = {
        "programa": extraer_valor_campo(bloque_info, "PROGRAMA", etiquetas),
        "area": extraer_valor_campo(bloque_info, "ÁREA", etiquetas) or extraer_valor_campo(bloque_info, "AREA", etiquetas),
        "titulo": extraer_valor_campo(bloque_info, "ASIGNATURA", etiquetas),
        "creditos": extraer_valor_campo(bloque_info, "CRÉDITOS", etiquetas) or extraer_valor_campo(bloque_info, "CREDITOS", etiquetas),
        "semestre": extraer_valor_campo(bloque_info, "SEMESTRE", etiquetas),
        "horas_directas": extraer_horas(ls, "HORAS PRESENCIA"),
        "horas_independientes": extraer_horas(ls, "HORAS DE TRABAJO"),
        "profesor": "",
        "correo": "",
    }
    return {k: limpiar_texto(v) for k, v in info.items()}


def es_numero_pagina(linea):
    return bool(re.fullmatch(r"\d+", linea.strip()))


def es_heading_interno(linea):
    n = norm(linea)
    return any(n.startswith(x) for x in [
        "RESULTADOS DE LA COMPETENCIA", "BIBLIOGRAFIA BASICA", "BIBLIOGRAFIA COMPLEMENTARIA",
        "RECURSOS DIGITALES", "INSTRUMENTOS DE EVALUACION", "CRITERIOS Y PORCENTAJES",
        "PORCENTAJES DE EVALUACION", "FALLAS Y LINEAMIENTOS"
    ])


def empieza_vineta(linea):
    l = linea.strip()
    if l.startswith(("•", "-", "", "▪", "■")):
        return True
    verbos = (
        "COMPRENDER", "ANALIZAR", "DISEÑAR", "INTERPRETAR", "APLICAR", "COMUNICAR",
        "ELABORAR", "EVALUAR", "FORMULAR", "SUSTENTAR", "PRESENTAR", "DOCUMENTAR",
        "IDENTIFICAR", "UTILIZAR", "CONSTRUIR", "EXPLICAR", "RESOLVER", "SELECCIONAR"
    )
    return norm(l).startswith(verbos)


def reflow_bloque(texto, activar_vinetas=True):
    lineas = []
    for l in (texto or "").splitlines():
        l = limpiar_texto(l)
        if not l or l.startswith("--- PAGINA") or es_numero_pagina(l):
            continue
        lineas.append(l)

    salida = []
    actual = ""
    actual_tipo = "normal"
    contexto_vinetas = False

    def flush():
        nonlocal actual, actual_tipo
        if actual.strip():
            salida.append((actual_tipo, actual.strip()))
        actual = ""
        actual_tipo = "normal"

    def append_to_actual(linea):
        nonlocal actual
        if actual.endswith("-"):
            actual = actual[:-1] + linea
        elif actual:
            actual += " " + linea
        else:
            actual = linea

    for l in lineas:
        n = norm(l)
        if es_heading_interno(l):
            flush()
            salida.append(("heading", l))
            contexto_vinetas = True
            continue
        if re.match(r"^\d+\.\s+", l):
            flush()
            actual = l
            actual_tipo = "normal"
            contexto_vinetas = n.endswith("IMPLICA") or n.endswith("INCLUYE")
            continue
        if activar_vinetas and (l.startswith(("•", "-", "", "▪", "■")) or (contexto_vinetas and empieza_vineta(l))):
            flush()
            actual = l.lstrip("•-▪■ ").strip()
            actual_tipo = "bullet"
            contexto_vinetas = True
            continue
        append_to_actual(l)
        if n.endswith("IMPLICA") or n.endswith("INCLUYE") or norm(actual).endswith("IMPLICA") or norm(actual).endswith("INCLUYE"):
            contexto_vinetas = True
    flush()
    return salida


def reflow_bibliografia(texto):
    out = []
    actual = ""

    def flush_ref():
        nonlocal actual
        if actual.strip():
            out.append(("bullet", actual.strip().lstrip("•- ")))
        actual = ""

    for raw in (texto or "").splitlines():
        linea = limpiar_texto(raw)
        if not linea or linea.startswith("--- PAGINA") or es_numero_pagina(linea):
            continue
        n = norm(linea)
        if n.startswith(("BIBLIOGRAFIA BASICA", "BIBLIOGRAFIA COMPLEMENTARIA", "RECURSOS DIGITALES")):
            flush_ref()
            out.append(("heading", linea))
            continue
        sin_num = re.sub(r"^\d+\.?\s+", "", linea).strip()
        es_nueva_ref = bool(re.search(r"\(\d{4}\)", sin_num[:180]))
        if es_nueva_ref:
            flush_ref()
            actual = sin_num
        else:
            if actual:
                actual += " " + linea
            else:
                actual = sin_num
    flush_ref()
    return out


def texto_sin_codigo(texto):
    texto = limpiar_texto(str(texto or ""))
    # Quita códigos visibles tipo CP1., CP2 -, RdAP1:, etc. cuando vengan al inicio.
    texto = re.sub(r"^\s*(?:CP\d+|RdAP\d+)\s*[\.\-:)]\s*", "", texto, flags=re.I).strip()
    return texto


def leer_matriz_alineacion(matriz_path, asignatura):
    """Lee una matriz plana. Hoja esperada: Matriz_Plana con columnas Curso, Tipo, Texto, Activo."""
    resultado = {"competencias_programa": [], "rdap_programa": [], "encontrado": False}
    if not matriz_path or not Path(matriz_path).exists() or load_workbook is None:
        return resultado
    wb = load_workbook(matriz_path, data_only=True)
    if "Matriz_Plana" not in wb.sheetnames:
        return resultado
    ws = wb["Matriz_Plana"]
    headers = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    hmap = {norm(h): i for i, h in enumerate(headers)}
    idx_curso = hmap.get("CURSO")
    idx_tipo = hmap.get("TIPO")
    idx_texto = hmap.get("TEXTO")
    idx_activo = hmap.get("ACTIVO")
    if idx_curso is None or idx_tipo is None or idx_texto is None:
        return resultado
    buscada = norm(asignatura)
    for row in ws.iter_rows(min_row=2, values_only=True):
        curso = row[idx_curso] if idx_curso < len(row) else None
        tipo = row[idx_tipo] if idx_tipo < len(row) else None
        texto = row[idx_texto] if idx_texto < len(row) else None
        activo = row[idx_activo] if idx_activo is not None and idx_activo < len(row) else "Sí"
        if not curso or not tipo or not texto:
            continue
        if norm(str(activo)) not in {"SI", "SÍ", "YES", "TRUE", "1", "ACTIVO"}:
            continue
        ncurso = norm(str(curso))
        if ncurso != buscada and not (buscada in ncurso or ncurso in buscada):
            continue
        resultado["encontrado"] = True
        item = ("bullet", texto_sin_codigo(texto))
        ntipo = norm(str(tipo))
        if ntipo in {"COMPETENCIA", "COMPETENCIAS", "CP"}:
            resultado["competencias_programa"].append(item)
        elif ntipo in {"RDA", "RDAP", "RESULTADO", "RESULTADOS", "RESULTADO DE APRENDIZAJE"}:
            resultado["rdap_programa"].append(item)
    # Elimina duplicados conservando orden
    for key in ["competencias_programa", "rdap_programa"]:
        seen = set()
        unique = []
        for tipo, texto in resultado[key]:
            k = norm(texto)
            if k and k not in seen:
                seen.add(k)
                unique.append((tipo, texto))
        resultado[key] = unique
    return resultado


def set_run_font(run, bold=False):
    run.font.name = FUENTE_CONTENIDO
    run.font.size = Pt(TAMANO_CONTENIDO)
    run.bold = bold


def set_paragraph_format(p, tipo="normal"):
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    if tipo == "bullet":
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-9)


def clear_paragraph(p):
    p.clear()
    set_paragraph_format(p)


def write_paragraph(p, texto, tipo="normal"):
    clear_paragraph(p)
    prefix = "• " if tipo == "bullet" else ""
    r = p.add_run(prefix + limpiar_texto(texto or ""))
    set_run_font(r, bold=(tipo == "heading"))
    set_paragraph_format(p, tipo)


def insert_paragraph_after(paragraph, texto="", tipo="normal"):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    write_paragraph(p, texto, tipo)
    return p


def buscar_parrafo(doc, contiene):
    objetivo = norm(contiene)
    for p in doc.paragraphs:
        if objetivo in norm(p.text):
            return p
    return None


def reemplazar_titulo(doc, titulo):
    p = buscar_parrafo(doc, "Título del curso")
    if p:
        write_paragraph(p, (titulo or "").upper(), "heading")


def limpiar_e_insertar_en_placeholder(doc, texto_placeholder, parrafos):
    p = buscar_parrafo(doc, texto_placeholder)
    if not p:
        return False
    if not parrafos:
        clear_paragraph(p)
        return True
    write_paragraph(p, parrafos[0][1], parrafos[0][0])
    ref = p
    for tipo, texto in parrafos[1:]:
        ref = insert_paragraph_after(ref, texto, tipo)
    return True


def escribir_celda(cell, texto):
    cell.text = limpiar_texto(texto or "")
    for p in cell.paragraphs:
        for r in p.runs:
            set_run_font(r)
        set_paragraph_format(p)


def diligenciar_tabla_inicial(doc, info, fecha_actualizacion):
    if not doc.tables:
        return
    t = doc.tables[0]
    try:
        escribir_celda(t.rows[0].cells[7], fecha_actualizacion)
    except Exception:
        pass
    try:
        escribir_celda(t.rows[1].cells[1], info.get("programa", ""))
    except Exception:
        pass
    try:
        escribir_celda(t.rows[4].cells[1], info.get("creditos", ""))
        escribir_celda(t.rows[4].cells[5], info.get("horas_directas", ""))
        escribir_celda(t.rows[4].cells[8], info.get("horas_independientes", ""))
    except Exception:
        pass


def limpiar_celda_texto(x):
    return limpiar_texto((x or "").replace("\n", " "))


def quitar_fecha_inicio(texto):
    return limpiar_texto(DATE_PREFIX_RE.sub("", texto or ""))


def es_fila_no_clase(sem, ses, con):
    n = norm(" ".join([sem or "", ses or "", con or ""]))
    if not n:
        return True
    if any(x in n for x in ["FESTIVO", "SEMANA SANTA", "EXAMEN FINAL", "PRESENTACION FINAL"]):
        return True
    # Evita encabezados repetidos.
    if "SEMANA" in n and "SESION" in n:
        return True
    return False


def extraer_tabla_tematicas_por_texto(bloque_tematicas):
    lines = [limpiar_texto(l) for l in (bloque_tematicas or "").splitlines() if limpiar_texto(l)]
    rows = []
    current_week = ""
    i = 0
    while i < len(lines):
        l = lines[i]
        n = norm(l)
        if n in {"SEMANA SESION FECHA", "SEMANA SESION FECHA CONTENIDO", "CONTENIDO"}:
            i += 1
            continue
        if any(x in n for x in ["FESTIVO", "SEMANA SANTA", "EXAMEN FINAL", "PRESENTACION FINAL"]):
            i += 1
            continue
        if re.fullmatch(r"\d{1,2}", l):
            # Puede ser semana o sesión. Si la siguiente línea es sesión, esta es semana.
            if i + 1 < len(lines) and re.fullmatch(r"\d{1,2}", lines[i + 1]):
                current_week = l
                i += 1
                continue
            # Si la siguiente línea es festivo/receso, esta línea sigue siendo semana.
            if i + 1 < len(lines) and any(x in norm(lines[i + 1]) for x in ["FESTIVO", "SEMANA SANTA"]):
                current_week = l
                i += 1
                continue
            # Si siguiente línea tiene fecha, esta línea es sesión.
            if i + 1 < len(lines) and DATE_PREFIX_RE.match(lines[i + 1]):
                ses = l
                con = quitar_fecha_inicio(lines[i + 1])
                # Junta continuaciones que no sean nuevos números/fechas/encabezados.
                j = i + 2
                while j < len(lines):
                    nj = norm(lines[j])
                    if re.fullmatch(r"\d{1,2}", lines[j]) or DATE_PREFIX_RE.match(lines[j]) or any(x in nj for x in ["FESTIVO", "SEMANA SANTA", "EVALUACION", "PRESENTACION FINAL", "EXAMEN FINAL"]):
                        break
                    con += " " + lines[j]
                    j += 1
                if not es_fila_no_clase(current_week, ses, con):
                    rows.append([current_week, ses, limpiar_texto(con)])
                i = j
                continue
        i += 1
    if rows:
        return [["Semana", "Sesión", "Contenido"]] + rows
    return None


def extraer_tabla_tematicas_pdf(path_pdf, bloque_tematicas=""):
    # Primero usa el texto plano: suele preservar mejor los espacios que pdfplumber en tablas extensas.
    por_texto = extraer_tabla_tematicas_por_texto(bloque_tematicas)
    if por_texto:
        return por_texto
    filas_finales = []
    with pdfplumber.open(path_pdf) as pdf:
        for page in pdf.pages:
            try:
                tablas = page.extract_tables()
            except Exception:
                tablas = []
            for tabla in tablas or []:
                if not tabla or len(tabla) < 2:
                    continue
                rows = [[limpiar_celda_texto(c) for c in row] for row in tabla if row and any(row)]
                texto = " ".join(" ".join(r) for r in rows)
                if not all(x in norm(texto) for x in ["SEMANA", "SESION", "CONTENIDO"]):
                    continue
                header_idx = None
                for i, r in enumerate(rows[:4]):
                    nr = [norm(c) for c in r]
                    if any("SEMANA" in c for c in nr) and any("SESION" in c for c in nr) and any("CONTENIDO" in c for c in nr):
                        header_idx = i
                        break
                if header_idx is None:
                    continue
                header = [norm(c) for c in rows[header_idx]]
                def idx_col(nombre):
                    for j, c in enumerate(header):
                        if nombre in c:
                            return j
                    return None
                i_sem = idx_col("SEMANA")
                i_ses = idx_col("SESION")
                i_con = idx_col("CONTENIDO")
                if i_sem is None or i_ses is None or i_con is None:
                    continue
                for r in rows[header_idx + 1:]:
                    sem = r[i_sem] if i_sem < len(r) else ""
                    ses = r[i_ses] if i_ses < len(r) else ""
                    con = quitar_fecha_inicio(r[i_con] if i_con < len(r) else "")
                    if not es_fila_no_clase(sem, ses, con):
                        filas_finales.append([sem, ses, con])
    if not filas_finales:
        return None
    return [["Semana", "Sesión", "Contenido"]] + filas_finales


def insertar_tabla_tematicas(doc, data):
    p = buscar_parrafo(doc, "Precise los contenidos o temáticas")
    if not p:
        return False
    clear_paragraph(p)
    if not data or len(data) <= 1:
        return False
    tabla = doc.add_table(rows=1, cols=3)
    tabla.style = "Table Grid"
    for j, val in enumerate(data[0][:3]):
        escribir_celda(tabla.cell(0, j), val)
    for fila in data[1:]:
        cells = tabla.add_row().cells
        for j in range(3):
            escribir_celda(cells[j], fila[j] if j < len(fila) else "")
    tbl_xml = tabla._tbl
    doc._body._body.remove(tbl_xml)
    p._p.addnext(tbl_xml)
    return True


def eliminar_parrafos_vacios_finales(doc, max_remove=80):
    removed = 0
    for p in reversed(doc.paragraphs):
        if removed >= max_remove:
            break
        if (p.text or "").strip():
            break
        el = p._element
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)
            removed += 1
    return removed


def nombre_salida(info, pdf_path):
    titulo = info.get("titulo", "") or Path(pdf_path).stem
    titulo = limpiar_texto(titulo)
    titulo = re.sub(r"[^A-Za-zÁÉÍÓÚáéíóúÑñ0-9_ -]+", "", titulo).strip()
    if not titulo:
        titulo = Path(pdf_path).stem
    return f"Microcurriculo {titulo}.docx"


def procesar_pdf(template_path, pdf_path, output_dir, matriz_path=None):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    texto = leer_pdf(pdf_path)
    bloques, encabezados = extraer_bloques(texto)
    info = extraer_info_general(bloques.get("informacion_general", ""))
    fecha_actualizacion = datetime.today().strftime("%d/%m/%Y")
    datos_matriz = leer_matriz_alineacion(matriz_path, info.get("titulo", "")) if matriz_path else {
        "competencias_programa": [], "rdap_programa": [], "encontrado": False
    }

    doc = Document(template_path)
    reemplazar_titulo(doc, info.get("titulo", ""))
    diligenciar_tabla_inicial(doc, info, fecha_actualizacion)

    limpiar_e_insertar_en_placeholder(doc, "Relacione las competencias del programa", datos_matriz.get("competencias_programa", []))
    limpiar_e_insertar_en_placeholder(doc, "Relacione el o los resultados de aprendizaje del programa", datos_matriz.get("rdap_programa", []))
    limpiar_e_insertar_en_placeholder(doc, "Especifique uno o máximo dos resultados de aprendizaje", reflow_bloque(bloques.get("resultados", ""), activar_vinetas=True))

    presentacion = bloques.get("presentacion", "")
    if bloques.get("continuidad", "").strip():
        presentacion = (presentacion.strip() + "\n\n" + bloques.get("continuidad", "").strip()).strip()
    limpiar_e_insertar_en_placeholder(doc, "Describa las características que definen la naturaleza del curso", reflow_bloque(presentacion, activar_vinetas=False))

    tabla_tematicas = extraer_tabla_tematicas_pdf(pdf_path, bloques.get("tematicas", ""))
    tematicas_insertadas = insertar_tabla_tematicas(doc, tabla_tematicas)

    limpiar_e_insertar_en_placeholder(doc, "Describa la tipología del curso", reflow_bloque(bloques.get("metodologia", ""), activar_vinetas=True))
    limpiar_e_insertar_en_placeholder(doc, "Defina los criterios con los cuáles serán evaluados", reflow_bloque(bloques.get("evaluacion", ""), activar_vinetas=True))
    limpiar_e_insertar_en_placeholder(doc, "Incluya la bibliografía básica", reflow_bibliografia(bloques.get("bibliografia", "")))

    parrafos_vacios_eliminados = eliminar_parrafos_vacios_finales(doc)

    salida = output_dir / nombre_salida(info, pdf_path)
    doc.save(salida)

    diagnostico = {
        "pdf": str(pdf_path),
        "salida": str(salida),
        "titulo": info.get("titulo", ""),
        "programa": info.get("programa", ""),
        "creditos": info.get("creditos", ""),
        "horas_directas": info.get("horas_directas", ""),
        "horas_independientes": info.get("horas_independientes", ""),
        "profesor": "NO DILIGENCIADO POR REGLA DEL PROYECTO",
        "correo": "NO DILIGENCIADO POR REGLA DEL PROYECTO",
        "fecha_actualizacion": fecha_actualizacion,
        "matriz_usada": bool(matriz_path),
        "asignatura_en_matriz": datos_matriz.get("encontrado", False),
        "competencias_programa_desde_matriz": len(datos_matriz.get("competencias_programa", [])),
        "rdap_desde_matriz": len(datos_matriz.get("rdap_programa", [])),
        "secciones_pdf": {k: bool((v or "").strip()) for k, v in bloques.items()},
        "filas_tematicas": len(tabla_tematicas) if tabla_tematicas else 0,
        "tematicas_insertadas": tematicas_insertadas,
        "nota_tematicas": "Tabla final Semana | Sesión | Contenido, sin columna Fecha. Se omiten festivos y recesos.",
        "parrafos_vacios_finales_eliminados": parrafos_vacios_eliminados,
    }
    return salida, diagnostico
